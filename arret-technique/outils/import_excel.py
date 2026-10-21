"""Convertit le planning Excel d'arrêt technique (.xlsm) en data.json.

Usage : python3 import_excel.py Planning.xlsm ../data.json
"""

import json
import re
import sys
import unicodedata
from datetime import date, datetime

import openpyxl

# Disposition du schéma : serpentin qui suit le sens du produit, puis les utilités.
LIGNE = [
    ["PARC A BILLETTES", "ENFOURNEMENT", "FOUR", "DEFOURNEMENT", "DÉCALAMINEUSE", "TRO-0", "SOUDEUSE", "TRO-1"],
    ["BOUCLEUSES", "VR2", "INTERMÉDIAIRE", "VR1", "BLOC CV1", "DEGROSSISSEUR", "EBAVUREUSE", "VR ESCAMOTABLE"],
    ["CURSEUR", "BLOC CV3", "BLOC CV4", "VR3", "BLOC CV5", "TEMPCORE", "LIGNE PETIT Ø", "TRANSFERT"],
    [None, None, None, None, None, None, "BOTTELAGE", "TABLE REFROIDISSOIRE"],
]
UTILITES = [
    "STATIONS EAU LAMINOIR", "STATION EAU n°1", "STATION EAU n°2", "REJET EAU INDUSTRIELLE",
    "LOCAL ADOUCISSEURS", "CIRCUIT MOTEURS", "CIRCUIT TEMPCORE", "CIRCUIT DIRECT LAMINOIR",
]
# Sens de passage du produit (flèches du schéma).
FLUX = [
    "PARC A BILLETTES", "ENFOURNEMENT", "FOUR", "DEFOURNEMENT", "DÉCALAMINEUSE", "TRO-0", "SOUDEUSE", "TRO-1",
    "VR ESCAMOTABLE", "EBAVUREUSE", "DEGROSSISSEUR", "BLOC CV1", "VR1", "INTERMÉDIAIRE", "VR2", "BOUCLEUSES",
    "CURSEUR", "BLOC CV3", "BLOC CV4", "VR3", "BLOC CV5", "TEMPCORE", "LIGNE PETIT Ø", "TRANSFERT",
    "TABLE REFROIDISSOIRE", "BOTTELAGE",
]
GENERAL = ["LAMINOIR", "ENFOURNEMENT - DEFOURNEMENT", "VR1 / VR2", "SOUS-STATION PRINCIPALE", "ATELIER"]

STATUTS = {"A FAIRE": "afaire", "EN COURS": "encours", "FAIT": "fait",
           "NON NECESSAIRE": "nonnecessaire", "REPORTÉE": "reportee"}
EQUIPES = [("mto", "MTO", 3), ("meca", "Entretien mécanique", 11),
           ("elec", "Entretien électrique", 19), ("resp", "Encadrement", 27)]


def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().upper()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", norm(s).lower()).strip("-")


def iso(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return None


def heures(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).lower().replace(",", ".").replace(" ", "")
    m = re.match(r"([\d.]+)(h|min)?", s)
    if not m:
        return None
    n = float(m.group(1))
    return round(n / 60, 2) if m.group(2) == "min" else n


def zone_layout():
    zones, x0, y0, w, h, gx, gy = [], 20, 40, 130, 70, 22, 70
    for r, row in enumerate(LIGNE):
        for c, nom in enumerate(row):
            if nom:
                zones.append({"id": slug(nom), "nom": nom, "type": "process",
                              "x": x0 + c * (w + gx), "y": y0 + r * (h + gy), "w": w, "h": h})
    y = y0 + len(LIGNE) * (h + gy) + 30
    for i, nom in enumerate(UTILITES):
        zones.append({"id": slug(nom), "nom": nom, "type": "utilite",
                      "x": x0 + i * (w + gx), "y": y, "w": w, "h": 56})
    y2 = y0 + 3 * (h + gy)
    for i, nom in enumerate(GENERAL):
        zones.append({"id": slug(nom), "nom": nom, "type": "general",
                      "x": x0 + i * (w + gx), "y": y2, "w": w, "h": h})
    return zones


def main(src, dst):
    wb = openpyxl.load_workbook(src, data_only=True)

    # --- Personnel -------------------------------------------------------
    ws = wb["bdd personnel"]
    personnel, vus = [], {}
    for eq_id, _, col in EQUIPES:
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=col + 1, max_col=col + 4, values_only=True):
            mat, nom, prenom, poste = row
            if not nom or not str(nom).strip():
                continue
            key = mat or norm(nom)
            if key in vus:
                continue
            p = {"id": "p%s" % (mat or slug(nom)), "matricule": mat, "nom": norm(nom),
                 "prenom": str(prenom or "").strip().title(), "poste": str(poste or "").strip(),
                 "equipe": eq_id, "absences": []}
            vus[key] = p
            personnel.append(p)

    def trouver(token):
        t = norm(token)
        if not t:
            return None
        exact = [p for p in personnel if t == p["nom"] or t in p["nom"].split(" ") or t == norm(p["prenom"])]
        if len(exact) == 1:
            return exact[0]
        pref = [p for p in personnel if len(t) >= 3 and (norm(p["prenom"]).startswith(t) or t.startswith(norm(p["prenom"])) and p["prenom"])]
        if len(pref) == 1:
            return pref[0]
        return None

    def intervenants(txt):
        ids = []
        for tok in re.split(r"[/,\-+&]", str(txt or "")):
            tok = tok.strip()
            if not tok:
                continue
            p = trouver(tok)
            if not p:
                pid = "x" + slug(tok)
                p = next((q for q in personnel if q["id"] == pid), None)
                if not p:
                    p = {"id": pid, "matricule": None, "nom": norm(tok), "prenom": "", "poste": "",
                         "equipe": "ext", "absences": []}
                    personnel.append(p)
            if p["id"] not in ids:
                ids.append(p["id"])
        return ids

    # --- Planning ----------------------------------------------------------
    ws = wb["Planning"]
    meta = {
        "titre": str(ws["A1"].value or "Arrêt technique").strip(),
        "debut": iso(ws["B2"].value), "fin": iso(ws["B3"].value),
        "finAffichage": max(iso(c.value) or "" for r in ws.iter_rows(min_row=8, max_row=11) for c in r) or iso(ws["B3"].value),
        "consignes": [l.strip() for l in str(ws["C2"].value or "").split("\n") if l.strip()],
        "jalons": [
            {"date": iso(ws["B4"].value), "libelle": "Sortie billettes four"},
            {"date": iso(ws["B5"].value), "libelle": "Rentrée billettes four"},
            {"date": iso(ws["B6"].value), "libelle": "Allumage four"},
        ],
        "capacite": {"1": 10, "2": 7, "3": 7, "4": 7, "5": 6.5, "6": 0, "0": 0},
        "tempsPerdu": 0.15,
        "modeCharge": "parIntervenant",
    }
    zones = zone_layout()
    zone_ids = {z["nom"]: z["id"] for z in zones}
    activites, n = [], 0
    for row in ws.iter_rows(min_row=10, max_row=ws.max_row, max_col=13, values_only=True):
        zone, service, desig, statut, interv, gmao, temps, moi, deb, duree, debr, dureer, _ = row
        if not desig:
            continue
        n += 1
        if zone not in zone_ids:
            z = {"id": slug(zone), "nom": zone, "type": "general", "x": 20, "y": 20, "w": 130, "h": 70}
            zones.append(z)
            zone_ids[zone] = z["id"]
        activites.append({
            "id": "a%03d" % n, "zone": zone_ids[zone], "service": service or "",
            "designation": re.sub(r"\s+", " ", str(desig)).strip(),
            "statut": STATUTS.get(statut, "afaire"), "intervenants": intervenants(interv),
            "gmao": str(gmao).replace("\n", " ") if gmao else "", "moi": moi or "",
            "dureeH": heures(duree) or heures(temps), "debut": iso(deb), "fin": iso(deb),
            "debutReel": iso(debr), "dureeReelleH": heures(dureer), "pont": "", "commentaire": "",
        })

    services = [r[0] for r in wb["Liste"].iter_rows(min_col=3, max_col=3, values_only=True) if r[0]]
    data = {
        "version": 1, "meta": meta,
        "equipes": [{"id": e, "nom": nom} for e, nom, _ in EQUIPES] + [{"id": "ext", "nom": "Entreprises / non référencés"}],
        "services": services, "zones": zones, "flux": [slug(z) for z in FLUX], "personnel": personnel, "activites": activites,
        "ponts": [{"id": "demag", "nom": "DEMAG", "zone": "Laminoir"},
                  {"id": "anselmi", "nom": "ANSELMI", "zone": ""}],
        "reservations": [],
    }
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print("%d activités, %d personnes, %d zones -> %s" % (len(activites), len(personnel), len(zones), dst))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
